file_utils
import os
import fitz  # PyMuPDF
from pptx import Presentation
from docx import Document
import logging

logger = logging.getLogger(__name__)

# File size limits (in bytes)
MAX_PDF_SIZE = 50 * 1024 * 1024  # 50MB

def extract_text_from_pdf(pdf_path):
    """Extracts text from PDF files - enhanced for large files"""
    logger.info(f"📄 Extracting text from PDF: {pdf_path}")
    try:
        # Check file size first
        if os.path.exists(pdf_path):
            file_size = os.path.getsize(pdf_path)
            if file_size > MAX_PDF_SIZE:
                file_size_mb = file_size / (1024 * 1024)
                raise Exception(f"PDF file too large ({file_size_mb:.1f}MB). Maximum supported size is {MAX_PDF_SIZE // (1024*1024)}MB.")
        
        if not os.path.exists(pdf_path):
            raise FileNotFoundError(f"File not found: {pdf_path}")

        doc = fitz.open(pdf_path)
        full_text = ""
        page_count = len(doc)
        
        # For very large files, process with progress tracking
        for page_num, page in enumerate(doc):
            page_text = page.get_text()
            full_text += page_text
            
            # Show progress for very large files
            if page_count > 50 and page_num % 10 == 0:
                logger.info(f"📄 Processed {page_num}/{page_count} pages...")
        
        doc.close()
        
        if not full_text.strip():
            logger.warning("⚠️ No text content found in PDF")
            raise Exception("No extractable text found in PDF. The file may contain only images or scanned content.")
            
        logger.info(f"✅ PDF extraction complete: {len(full_text)} total characters")
        return full_text.strip()
        
    except MemoryError:
        logger.error("❌ Memory error processing large PDF")
        raise Exception("PDF is too large to process and caused memory issues. Please try a smaller file.")
    except Exception as e:
        logger.error(f"❌ PDF extraction failed: {e}")
        # Re-raise with more context if it's a size issue
        if "too large" in str(e).lower():
            raise
        raise Exception(f"Failed to extract text from PDF: {str(e)}")

def extract_text_from_pdf_bytes(pdf_bytes):
    """Extract text from PDF bytes without temp files"""
    try:
        # Check size of bytes
        if len(pdf_bytes) > MAX_PDF_SIZE:
            file_size_mb = len(pdf_bytes) / (1024 * 1024)
            raise Exception(f"PDF file too large ({file_size_mb:.1f}MB). Maximum supported size is {MAX_PDF_SIZE // (1024*1024)}MB.")
            
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        full_text = ""
        for page in doc:
            full_text += page.get_text()
        doc.close()
        
        if not full_text.strip():
            raise Exception("No extractable text found in PDF.")
            
        return full_text.strip()
    except MemoryError:
        logger.error("❌ Memory error processing large PDF bytes")
        raise Exception("PDF is too large to process. Please try a smaller file.")
    except Exception as e:
        logger.error(f"❌ PDF bytes extraction failed: {e}")
        raise

def extract_text_from_pptx(pptx_path):
    """Extracts text from PowerPoint (PPTX) files."""
    logger.info(f"📊 Extracting text from PPTX: {pptx_path}")
    try:
        # Check file size
        if os.path.exists(pptx_path):
            file_size = os.path.getsize(pptx_path)
            if file_size > 50 * 1024 * 1024:  # 50MB
                file_size_mb = file_size / (1024 * 1024)
                raise Exception(f"PPTX file too large ({file_size_mb:.1f}MB). Maximum supported size is 50MB.")

        if not os.path.exists(pptx_path):
            raise FileNotFoundError(f"File not found: {pptx_path}")

        prs = Presentation(pptx_path)
        full_text = ""
        for slide_num, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    full_text += shape.text + "\n"
            logger.debug(f"📊 Slide {slide_num + 1} processed")
        
        if not full_text.strip():
            raise Exception("No text content found in PowerPoint file.")
            
        logger.info(f"✅ PPTX extraction complete: {len(full_text)} total characters")
        return full_text.strip()
    except MemoryError:
        logger.error("❌ Memory error processing large PPTX")
        raise Exception("PowerPoint file is too large to process. Please try a smaller file.")
    except Exception as e:
        logger.error(f"❌ PPTX extraction failed: {e}")
        raise Exception(f"Failed to extract text from PowerPoint: {str(e)}")

def extract_text_from_docx(docx_path):
    """Extracts text from Word (DOCX) files."""
    logger.info(f"📝 Extracting text from DOCX: {docx_path}")
    try:
        # Check file size
        if os.path.exists(docx_path):
            file_size = os.path.getsize(docx_path)
            if file_size > 50 * 1024 * 1024:  # 50MB
                file_size_mb = file_size / (1024 * 1024)
                raise Exception(f"DOCX file too large ({file_size_mb:.1f}MB). Maximum supported size is 50MB.")

        if not os.path.exists(docx_path):
            raise FileNotFoundError(f"File not found: {docx_path}")

        doc = Document(docx_path)
        full_text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        
        if not full_text.strip():
            raise Exception("No text content found in Word document.")
            
        logger.info(f"✅ DOCX extraction complete: {len(full_text)} total characters")
        return full_text.strip()
    except MemoryError:
        logger.error("❌ Memory error processing large DOCX")
        raise Exception("Word document is too large to process. Please try a smaller file.")
    except Exception as e:
        logger.error(f"❌ DOCX extraction failed: {e}")
        raise Exception(f"Failed to extract text from Word document: {str(e)}")